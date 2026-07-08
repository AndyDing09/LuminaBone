"""
Assemble the per-location triangulation slide PNGs into the two deck files:

    BONE_triangulation_slides.pptx   (one full-slide image per location)
    BONE_triangulation_slides.pdf    (same pages as a PDF)

Run AFTER make_triangulation_slides.py:
    py make_triangulation_slides.py
    py make_slides_deck.py
"""

import glob
import os

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SLIDES_DIR = os.path.join(PROJECT_ROOT, "depth_outputs", "triangulation_slides")
PPTX_OUT = os.path.join(PROJECT_ROOT, "presentations", "BONE_triangulation_slides.pptx")
PDF_OUT = os.path.join(PROJECT_ROOT, "presentations", "BONE_triangulation_slides.pdf")

# 13.33 x 7.5 in widescreen, same as the original deck.
SLIDE_W = Emu(12191695)
SLIDE_H = Emu(6858000)


def main():
    pngs = sorted(glob.glob(os.path.join(SLIDES_DIR, "loc*_triangulation.png")))
    if not pngs:
        raise SystemExit(f"no slide PNGs in {SLIDES_DIR} - run make_triangulation_slides.py first")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    for png in pngs:
        slide = prs.slides.add_slide(blank)
        with Image.open(png) as im:
            w, h = im.size
        # Fit the image inside the slide, centered, preserving aspect.
        scale = min(SLIDE_W / w, SLIDE_H / h)
        pw, ph = int(w * scale), int(h * scale)
        slide.shapes.add_picture(png, Emu((SLIDE_W - pw) // 2),
                                 Emu((SLIDE_H - ph) // 2), Emu(pw), Emu(ph))
    prs.save(PPTX_OUT)
    print(f"{len(pngs)} slides -> {PPTX_OUT}")

    pages = [Image.open(p).convert("RGB") for p in pngs]
    pages[0].save(PDF_OUT, save_all=True, append_images=pages[1:],
                  resolution=115.0)
    for p in pages:
        p.close()
    print(f"{len(pages)} pages -> {PDF_OUT}")


if __name__ == "__main__":
    main()
